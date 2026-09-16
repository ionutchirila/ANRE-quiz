"""Export questions from the ANRE Excel workbook into questions.js."""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "Chestionare ANRE IV A+B Quiz(1).xlsx"
OUT = ROOT / "questions.js"


def drawing_excel_rows(xlsx: Path) -> set[int]:
    with zipfile.ZipFile(xlsx) as z:
        xml = z.read("xl/drawings/drawing1.xml").decode("utf-8")
    rows = re.findall(
        r"<xdr:from>\s*<xdr:col>\d+</xdr:col>\s*<xdr:colOff>\d+</xdr:colOff>\s*<xdr:row>(\d+)</xdr:row>",
        xml,
    )
    return {int(r) + 1 for r in rows}


def letters_from_ijk(i, j, k) -> list[str]:
    found = []
    for value, letter in ((i, "a"), (j, "b"), (k, "c")):
        text = "" if value is None else str(value).strip().lower()
        if text == letter:
            found.append(letter)
    return found


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=False)
    ws = wb["Sheet1"]
    ws2 = wb["Sheet2"]
    corrections = {}
    for row in ws2.iter_rows(min_row=3, max_row=7, values_only=True):
        corrections[int(row[0])] = str(row[7]).strip().lower()

    draw_rows = drawing_excel_rows(XLSX)
    questions = []
    skipped = []

    for excel_row in range(7, 868):
        row = ws[excel_row]
        try:
            qid = int(row[0].value)
        except (TypeError, ValueError):
            continue

        text = (row[1].value or "").strip()
        opts = []
        for col in (2, 3, 4):
            value = row[col].value
            opts.append("" if value is None else str(value).strip())

        found = letters_from_ijk(row[8].value, row[9].value, row[10].value)
        if qid in corrections:
            answers = [corrections[qid]]
        elif found:
            answers = found
        else:
            skipped.append(qid)
            continue

        questions.append(
            {
                "id": qid,
                "text": text,
                "options": {"a": opts[0], "b": opts[1], "c": opts[2]},
                "answers": answers,
                "hasDiagram": excel_row in draw_rows,
                "complete": all(opts),
            }
        )

    payload = {
        "title": "Chestionar ANRE IV A+B",
        "source": "Întrebări septembrie 2024",
        "category": "Electrotehnică, mașini și rețele electrice",
        "skippedWithoutAnswer": skipped,
        "questions": questions,
    }
    OUT.write_text(
        "window.QUIZ_DATA = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )
    print(f"Wrote {len(questions)} questions to {OUT}")


if __name__ == "__main__":
    main()
